import openpyxl
import glob

for file in glob.glob("out/*.xlsx"):
    wb = openpyxl.load_workbook(file, read_only=True)
    print(f"File: {file.split('/')[-1]}")
    print(f"Tabs: {wb.sheetnames}")
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        headers = []
        for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
            headers = [str(h) for h in row if h]
        print(f"  Tab '{sheet}' Headers: {headers}")
    print("-" * 40)
